import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from docpipe import convert_record
from docpipe import converters, manifest


class XlsxTests(unittest.TestCase):
    def _book(self, root, build):
        path = Path(root) / "sample.xlsx"
        workbook = Workbook()
        build(workbook)
        workbook.save(path)
        workbook.close()
        return path

    def test_values_stay_values_and_hidden_sheet_warns(self):
        def build(workbook):
            sheet = workbook.active
            sheet.title = "표"
            sheet.append(["연도", "값", "빈칸"])
            sheet.append([2025, 10, None])
            hidden = workbook.create_sheet("RAW")
            hidden.sheet_state = "hidden"
            hidden["A1"] = "secret"

        with tempfile.TemporaryDirectory() as td:
            result = convert_record(self._book(td, build))

        self.assertEqual(result["method"], "XLSX 구조 추출 (openpyxl)")
        self.assertIn("| 2025 | 10 |", result["text"])
        self.assertNotIn("NaN", result["text"])
        self.assertIn("RAW (원본에서 숨김 시트)", result["text"])
        self.assertIn("RAW", result["warning"])

    def test_vertical_merge_fills_but_horizontal_merge_does_not_duplicate(self):
        def build(workbook):
            sheet = workbook.active
            sheet.merge_cells("A1:A2")
            sheet["A1"] = "세로"
            sheet.merge_cells("B1:C1")
            sheet["B1"] = "가로"

        with tempfile.TemporaryDirectory() as td:
            sheets, _ = converters.convert_with_xlsx(self._book(td, build))

        text = sheets[0]
        self.assertGreaterEqual(text.count("세로"), 2)
        self.assertEqual(text.count("가로"), 1)

    def test_ragged_rows_use_widest_row(self):
        rendered = converters._rows_markdown([["제목"], ["a", "b", "c"]])
        self.assertEqual(
            rendered,
            "| 제목 |  |  |\n|---|---|---|\n| a | b | c |",
        )


class ManifestTests(unittest.TestCase):
    def test_skip_does_not_replace_a_richer_existing_record(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / manifest.MANIFEST_NAME
            manifest.write_manifest(
                path,
                [manifest.entry("a.pdf", "converted", method="OCR", warning="check")],
            )
            document = manifest.write_manifest(
                path,
                [
                    manifest.entry("a.pdf", "skipped"),
                    manifest.entry("b.docx", "skipped"),
                ],
            )

        rows = {row["path"]: row for row in document["files"]}
        self.assertEqual(rows["a.pdf"]["outcome"], "converted")
        self.assertEqual(rows["a.pdf"]["method"], "OCR")
        self.assertEqual(rows["b.docx"]["outcome"], "skipped")
        self.assertEqual(document["counts"]["warned"], 1)

    def test_broken_manifest_is_rebuilt(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / manifest.MANIFEST_NAME
            path.write_text("{ broken", encoding="utf-8")
            document = manifest.write_manifest(
                path, [manifest.entry("a.pdf", "converted")]
            )
            parsed = json.loads(path.read_text(encoding="utf-8"))

        self.assertEqual(len(document["files"]), 1)
        self.assertEqual(parsed["schema"], 1)


if __name__ == "__main__":
    unittest.main()
